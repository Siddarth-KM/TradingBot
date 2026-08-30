#!/usr/bin/env python3
"""Agent 2: Excel Writer. Reads signals from VM via SSH, appends week block to InvestmentTracker.xlsx."""

import argparse, json, sys, shutil, subprocess, logging, time
from datetime import datetime, timedelta
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Border, Side, PatternFill

EXCEL_PATH    = r"C:\Users\sidda\OneDrive\Documents\InvestmentTracker.xlsx"
LOG           = Path("logs/excel.log")
TRACKER       = Path("agents/last_excel_week.txt")
SSH_KEY       = r"C:\Users\sidda\Downloads\oracle-tradingbot.key"
VM            = "ubuntu@137.131.26.201"
REMOTE        = "/opt/tradingbot/signals/current_signals.json"
RETRY_SECS    = 600   # retry interval if Excel is locked
WAIT_POLL_SEC = 300   # re-check VM every 5 min waiting for new signals
WAIT_MAX_MIN  = 120   # give up after 2 hours

logging.basicConfig(filename=str(LOG), level=logging.INFO,
                    format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger()
log.addHandler(logging.StreamHandler(sys.stdout))

def fetch_signals():
    result = subprocess.run(
        ["ssh", "-i", SSH_KEY, "-o", "StrictHostKeyChecking=no",
         "-o", "BatchMode=yes", "-o", "ConnectTimeout=15",
         VM, f"cat {REMOTE}"],
        capture_output=True, text=True, timeout=30
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip())
    return json.loads(result.stdout)

# ── Step 1: SSH read signals — wait for new data if not ready yet ─────────────
# Normal weekly operation needs no flags: signals come from the VM and the week
# number follows the email agent. --signals-file replays an archived signals
# JSON to backfill a missed week; --week pins the week number for when
# last_sent_week.txt is not the right source (e.g. no email was sent).
ap = argparse.ArgumentParser(description="Agent 2: write a week block to InvestmentTracker.xlsx")
ap.add_argument("--signals-file", help="read signals from this local JSON instead of the VM")
ap.add_argument("--week", type=int, help="week number to write (default: agents/last_sent_week.txt)")
args = ap.parse_args()

if args.signals_file:
    try:
        data = json.loads(Path(args.signals_file).read_text())
    except Exception as e:
        log.error(f"Could not read {args.signals_file}: {e}")
        sys.exit(1)
    log.info(f"REPLAY: signals from {args.signals_file} (VM not contacted)")
else:
    try:
        data = fetch_signals()
    except Exception as e:
        log.error(f"SSH/parse failed: {e}")
        sys.exit(1)

generated_at = data["timestamp"]
signal_ts    = datetime.fromisoformat(generated_at)

if TRACKER.exists() and not args.signals_file:
    stored     = json.loads(TRACKER.read_text())
    stored_ts  = datetime.fromisoformat(stored["timestamp"])
    wait_start = datetime.now()
    while signal_ts <= stored_ts:
        elapsed = (datetime.now() - wait_start).total_seconds() / 60
        if elapsed >= WAIT_MAX_MIN:
            log.info(f"No new signals after {WAIT_MAX_MIN}min — exiting.")
            sys.exit(0)
        log.info(f"Signals not yet updated (still week {stored['week']}). Waiting 5 min...")
        time.sleep(WAIT_POLL_SEC)
        try:
            data = fetch_signals()
            generated_at = data["timestamp"]
            signal_ts    = datetime.fromisoformat(generated_at)
        except Exception as e:
            log.warning(f"Refetch failed: {e}")

seen, signals = set(), []
for s in data["all_signals"]:
    t = s["ticker"]
    if t not in seen:
        seen.add(t)
        signals.append({
            "ticker":      t,
            "pred":        s["predicted_return"],
            "probability": s["direction_probability"],
            "buy_price":   s["last_close"],
            "sell_limit":  s["limit_sell"],
        })
signals.sort(key=lambda x: x["pred"], reverse=True)
log.info(f"Loaded {len(signals)} signals (generated_at={generated_at})")

# ── Derive week number and date range ─────────────────────────────────────────
# Use email tracker for week number continuity
if args.week is not None:
    week_number = args.week
else:
    email_tracker = json.loads(Path("agents/last_sent_week.txt").read_text())
    week_number   = email_tracker["week"]
monday           = signal_ts - timedelta(days=signal_ts.weekday())
tuesday          = monday + timedelta(days=1)
following_monday = monday + timedelta(days=7)
date_range = f"{tuesday.month}/{tuesday.day} - {following_monday.month}/{following_monday.day}"
log.info(f"Week {week_number} | {date_range}")

# ── Step 3–5: Write to Excel with retry on PermissionError ───────────────────
def write_excel():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # Find starting capital from last Total row
    starting_capital = None
    for row in ws.iter_rows():
        a_val, k_val = row[0].value, row[10].value
        if isinstance(a_val, str) and a_val.strip() == "Total" and k_val not in (None, ""):
            try:
                starting_capital = float(k_val)
            except (TypeError, ValueError):
                pass
    if starting_capital is None:
        log.error("No prior Total row with non-blank Final Amount — aborting.")
        sys.exit(1)
    log.info(f"Starting capital: ${starting_capital:,.2f}")

    # Backup
    backup = EXCEL_PATH.replace(".xlsx", f"_backup_{datetime.today().strftime('%Y%m%d')}.xlsx")
    shutil.copy2(EXCEL_PATH, backup)

    # Find first empty row (with one blank gap after last content)
    last_content = 0
    for r in range(ws.max_row, 0, -1):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, 12)):
            last_content = r
            break
    first_empty = last_content + 2
    data_start  = first_empty + 2

    # Week header
    ws.cell(row=first_empty, column=1).value = f"Week {week_number}"
    ws.cell(row=first_empty, column=2).value = date_range
    ws.cell(row=first_empty, column=1).font  = Font(bold=True)

    # Column headers
    for ci, hdr in enumerate(["Ticker","Timeframe","Pred","Probability","Buy Price",
                               "Sell Limit","Final Sell","Percent Return",
                               "Shares Purchased","Start Amount","Final Amount"], 1):
        ws.cell(row=first_empty+1, column=ci).value = hdr
        ws.cell(row=first_empty+1, column=ci).font  = Font(bold=True)

    PCT  = "0.00%"
    CURR = '"$"#,##0.00'

    # Data rows
    for i, sig in enumerate(signals):
        r = data_start + i
        ws.cell(row=r, column=1).value  = sig["ticker"]
        ws.cell(row=r, column=2).value  = 5

        c3 = ws.cell(row=r, column=3);  c3.value = sig["pred"];                  c3.number_format = PCT
        c4 = ws.cell(row=r, column=4);  c4.value = sig["probability"] / 100;     c4.number_format = PCT
        c5 = ws.cell(row=r, column=5);  c5.value = sig["buy_price"];             c5.number_format = CURR
        c6 = ws.cell(row=r, column=6);  c6.value = sig["sell_limit"];            c6.number_format = CURR
        c7 = ws.cell(row=r, column=7);                                           c7.number_format = CURR      # Final Sell — user fills
        c8 = ws.cell(row=r, column=8);  c8.value = f'=IF(G{r}="","",(G{r}-E{r})/E{r})'; c8.number_format = PCT
        c9 = ws.cell(row=r, column=9);                                           c9.number_format = "0.00"    # Shares Purchased — user fills
        c10 = ws.cell(row=r, column=10); c10.number_format = CURR
        c11 = ws.cell(row=r, column=11); c11.value = f'=IF(OR(I{r}="",G{r}=""),"",I{r}*G{r})'; c11.number_format = CURR

    total_r = data_start + len(signals)
    ws.cell(row=total_r,   column=1).value = "Total";             ws.cell(row=total_r,   column=1).font = Font(bold=True)
    ws.cell(row=total_r,   column=2).value = 5
    ws.cell(row=total_r,   column=7).number_format  = CURR
    ws.cell(row=total_r,   column=8).number_format  = PCT
    ws.cell(row=total_r,   column=10).number_format = CURR
    ws.cell(row=total_r,   column=11).number_format = CURR
    ws.cell(row=total_r+1, column=1).value = "S&P 500(Adjusted)"; ws.cell(row=total_r+1, column=2).value = 5
    ws.cell(row=total_r+1, column=7).number_format  = CURR
    ws.cell(row=total_r+1, column=8).number_format  = PCT
    ws.cell(row=total_r+1, column=10).number_format = CURR
    ws.cell(row=total_r+1, column=11).number_format = CURR
    ws.cell(row=total_r+2, column=1).value = "S&P 500";           ws.cell(row=total_r+2, column=2).value = 5
    ws.cell(row=total_r+2, column=7).number_format  = CURR
    ws.cell(row=total_r+2, column=8).number_format  = PCT
    ws.cell(row=total_r+2, column=10).number_format = CURR
    ws.cell(row=total_r+2, column=11).number_format = CURR

    for mi, (la, lc) in enumerate([
        ("Cum. Return","YTD Cum. Return"), ("Sharpe Ratio (Cum.)","YTD Sharpe Ratio"),
        ("Max Drawdown (Cum.)","YTD Max Drawdown"), ("Calmar Ratio (Cum.)","YTD Calmar Ratio"),
        ("Sortino Ratio (Cum.)","YTD Sortino Ratio"), ("Info. Ratio (Cum.)","YTD Info. Ratio"),
        ("Win Rate","YTD Win Rate"), ("Alpha vs SPY (Cum.)","YTD Alpha vs SPY"),
        ("Payoff Ratio (Cum.)","Payoff Ratio (YTD)"), ("Profit Factor (Cum.)","Profit Factor (YTD)"),
    ]):
        mr = total_r + 3 + mi
        ws.cell(row=mr, column=1).value = la
        ws.cell(row=mr, column=3).value = lc

    # ── Step 5b: Formatting ───────────────────────────────────────────────────
    T         = Side(border_style="thin")
    M         = Side(border_style="medium")
    BLACK     = "FF000000"
    blue_fill = PatternFill(patternType="solid", fgColor="DEEAF1")
    n         = len(signals)

    # Column widths — match Week 37 template exactly
    widths = {1:22, 2:12, 3:10, 4:14, 5:12, 6:12, 7:12, 8:17, 9:19, 10:14, 11:14}
    for col_idx, w in widths.items():
        ws.column_dimensions[chr(64 + col_idx)].width = w

    def apply_row(row, top_side=T, bold_a=False):
        bold   = Font(bold=True, color=BLACK)
        normal = Font(color=BLACK)
        for col in range(1, 12):
            ws.cell(row=row, column=col).border = Border(
                top=top_side, bottom=T,
                left=(M if col == 1 else T),
                right=T,
            )
            ws.cell(row=row, column=col).font = bold if (col == 1 or bold_a) else normal

    # Week header row — bold only, no border
    for col in range(1, 12):
        ws.cell(row=first_empty, column=col).font = Font(bold=True, color=BLACK)

    # Column header row — medium top, all bold
    apply_row(first_empty + 1, top_side=M, bold_a=True)

    # Data rows — thin top, col A bold
    for i in range(n):
        apply_row(data_start + i, top_side=T, bold_a=False)

    # Total row — medium top, col A bold
    apply_row(total_r, top_side=M, bold_a=False)

    # S&P 500 rows — thin top, col A bold
    apply_row(total_r + 1, top_side=T, bold_a=False)
    apply_row(total_r + 2, top_side=T, bold_a=False)

    # Metric label rows — blue fill on col A, bold labels, medium top on first row
    r_metric0 = total_r + 3
    for i in range(10):
        r = r_metric0 + i
        ws.cell(r, 1).fill   = blue_fill
        ws.cell(r, 1).font   = Font(bold=True, color=BLACK)
        ws.cell(r, 3).font   = Font(bold=True, color=BLACK)
        if i == 0:
            ws.cell(r, 1).border = Border(top=M)

    wb.save(EXCEL_PATH)

attempt = 1
while True:
    try:
        log.info(f"Attempt {attempt}: writing Week {week_number} to Excel...")
        write_excel()
        log.info(f"Week {week_number} written OK | {len(signals)} tickers | {date_range}")
        break
    except PermissionError:
        log.warning(f"Attempt {attempt} failed — file locked. Retrying in 10 min.")
        attempt += 1
        time.sleep(RETRY_SECS)

# ── Update tracker ────────────────────────────────────────────────────────────
TRACKER.write_text(json.dumps({"week": week_number, "timestamp": generated_at}))
